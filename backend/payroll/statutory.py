"""EPF and ETF statutory outputs for one payroll month.

EPF e-Return (EPF Department "EPF e-Return System" instructions):
  * EVEMC.TXT – Contribution Detail File, one line per member
  * EVEMP.TXT – Payment Summary File, one line
  Plain ASCII, fixed-width fields. Numeric "9.2" = 10 characters, e.g. 0001535.73;
  "11.2" = 12 characters. The EPF Department checks an employer's first files before
  live submission, so test the first upload with EPF / your bank.

ETF: the ETF Board's member-file layout is issued separately to larger employers, so the
ETF outputs here are contribution schedules (R4 / R1 support and the half-yearly Form II).
"""

import calendar
from html import escape
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import PayrollPeriod, Payslip
from workforce.reports import safe_cell

PAY_MODES = [("1", "Cheque"), ("2", "Cash"), ("3", "Money order"), ("4", "Direct debit / online")]


# ---------------------------------------------------------------- fixed-width helpers
def txt(v, n):
    return str(v or "").upper()[:n].ljust(n)


def num(v, n):
    s = "".join(ch for ch in str(v or "") if ch.isdigit())
    return s[-n:].rjust(n, "0")


def dec(v, int_digits):
    """e.g. dec(1535.73, 7) -> '0001535.73' (9.2 field)."""
    v = Decimal(v or 0).quantize(Decimal("0.01"))
    s = f"{v:.2f}"
    return s.rjust(int_digits + 3, "0")[-(int_digits + 3):]


# ---------------------------------------------------------------- data
@dataclass
class Member:
    slip: Payslip
    nic: str
    surname: str
    initials: str
    member_no: str
    earnings: Decimal
    employee: Decimal
    employer: Decimal
    total: Decimal
    etf: Decimal
    status: str
    days: Decimal
    grade: str

    @property
    def name(self):
        return f"{self.initials} {self.surname}".strip()


def members(period):
    """EPF members for the month, from processed payslips."""
    out = []
    y, m = period.year, period.month
    days_in_month = Decimal(period.frozen_company.no_pay_divisor)
    for s in period.payslips.select_related("employee") \
            .order_by("employee__epf_no", "employee__emp_no"):
        e = s.frozen_employee
        if not getattr(e, "epf_eligible", False): continue
        if s.epf_base <= 0:
            continue
        status = "E"
        if e.join_date and (e.join_date.year, e.join_date.month) == (y, m):
            status = "N"
        if e.leave_date and (e.leave_date.year, e.leave_date.month) == (y, m):
            status = "V"
        inp = period.inputs.filter(employee_id=s.employee_id).first()
        no_pay = inp.no_pay_days if inp else Decimal(0)
        surname, initials = e.epf_name()
        out.append(Member(
            slip=s, nic=e.nic, surname=surname, initials=initials, member_no=e.epf_no,
            earnings=s.epf_base, employee=s.epf_employee, employer=s.epf_employer,
            total=s.epf_employee + s.epf_employer, etf=s.etf_employer, status=status,
            days=Decimal(s.snapshot.get("days", str(max(days_in_month - no_pay, Decimal(0))))), grade=e.epf_occupation_grade,
        ))
    out.sort(key=lambda m: (int(m.member_no) if str(m.member_no).isdigit() else 10**6, m.slip.employee.emp_no))
    return out


def problems(period, mems=None):
    """Things that must be fixed before the EPF files will be accepted."""
    c = period.frozen_company
    issues = []
    if not c.epf_zone or not c.epf_employer_no:
        issues.append("Company EPF zone code and employer number are not set (Settings → Companies).")
    for mbr in mems if mems is not None else members(period):
        e = mbr.slip.frozen_employee
        missing = [label for label, v in [("NIC", e.nic), ("EPF member no.", e.epf_no)] if not v]
        if missing:
            issues.append(f"{e.full_name}: missing {', '.join(missing)}")
    if not period.payslips.exists():
        issues.append("Payroll has not been processed for this month yet.")
    return issues


def period_code(period):
    return f"{period.year}{period.month:02d}"


# ---------------------------------------------------------------- EPF e-return text files
def evemc(period, submission=1):
    c = period.frozen_company
    lines = []
    for mbr in members(period):
        lines.append("".join([
            txt(mbr.nic, 20), txt(mbr.surname, 40), txt(mbr.initials, 20), num(mbr.member_no, 6),
            dec(mbr.total, 7), dec(mbr.employer, 7), dec(mbr.employee, 7), dec(mbr.earnings, 9),
            txt(mbr.status, 1), txt(c.epf_zone, 1), num(c.epf_employer_no, 6), period_code(period),
            num(submission, 2), dec(mbr.days, 4), num(mbr.grade, 3) if mbr.grade else "   ",
        ]))
    return "\r\n".join(lines) + "\r\n"


def evemp(period, submission=1, pay_mode="4", pay_ref="", pay_date=None):
    c = period.frozen_company
    mems = members(period)
    total = sum((m.total for m in mems), Decimal(0))
    pay_date = pay_date or date.today()
    line = "".join([
        txt(c.epf_zone, 1), num(c.epf_employer_no, 6), period_code(period), num(submission, 2),
        dec(total, 9), num(len(mems), 5), num(pay_mode, 1), txt(pay_ref, 20),
        pay_date.strftime("%Y%m%d"), txt(c.epf_district_office, 2),
    ])
    return line + "\r\n"


def epf_zip(period, **kw):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("EVEMC.TXT", evemc(period, kw.get("submission", 1)).encode("ascii", "replace"))
        z.writestr("EVEMP.TXT", evemp(period, **kw).encode("ascii", "replace"))
        z.writestr("README.txt", (
            f"EPF e-Return for {period.frozen_company.name} – {period.label}\r\n\r\n"
            "EVEMC.TXT  Contribution Detail File (one line per member)\r\n"
            "EVEMP.TXT  Payment Summary File\r\n\r\n"
            "Upload both files through your bank's EPF e-Return facility or as instructed by the EPF Department.\r\n"
            "The EPF Department checks a new employer's first files before live submission.\r\n"
            "Prepared by Jay's HR.\r\n"))
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def _sheet(ws, title, sub, headers, rows, money_cols, brand):
    ws["A1"], ws["A2"] = title, sub
    ws["A1"].font = Font(size=13, bold=True)
    ws["A2"].font = Font(size=10, color="555555")
    hr = 4
    for i, h in enumerate(headers, 1):
        cell = ws.cell(hr, i, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=brand)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r_i, row in enumerate(rows, hr + 1):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, float(v) if isinstance(v, Decimal) else safe_cell(v) if isinstance(v, str) else v)
            if c_i in money_cols:
                cell.number_format = "#,##0.00"
    tot = hr + len(rows) + 1
    ws.cell(tot, 2, "TOTAL").font = Font(bold=True)
    for c_i in money_cols:
        col = get_column_letter(c_i)
        cell = ws.cell(tot, c_i, f"=SUM({col}{hr + 1}:{col}{tot - 1})" if rows else 0)
        cell.font, cell.number_format = Font(bold=True), "#,##0.00"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = ws.cell(hr + 1, 1)
    return tot


def epf_excel(period, submission=1):
    """Contribution Detail in the same field order as the EPF template (for checking or the EPF macro)."""
    c = period.frozen_company
    wb = Workbook()
    ws = wb.active
    ws.title = "Contribution Detail"
    mems = members(period)
    headers = ["NIC Number", "Name (initials + surname)", "Last Name", "Initials", "Member No",
               "Total Contribution", "Employer Contribution (12%)", "Member Contribution (8%)",
               "Total Earnings", "Member Status (E/N/V)", "Zone", "Employer No", "Contribution Period",
               "Submission ID", "Days Worked", "Occupation Grade"]
    rows = [[m.nic, m.name, m.surname, m.initials, m.member_no, m.total, m.employer, m.employee, m.earnings,
             m.status, c.epf_zone, c.epf_employer_no, period_code(period), submission, m.days, m.grade]
            for m in mems]
    _sheet(ws, f"{c.name} – EPF contributions {period.label}",
           f"EPF employer no. {c.epf_zone}/{c.epf_employer_no or '—'} · {len(mems)} members · Prepared by Jay's HR",
           headers, rows, {6, 7, 8, 9}, c.theme()["brand"].lstrip("#"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def etf_excel(period):
    c = period.frozen_company
    wb = Workbook()
    ws = wb.active
    ws.title = "ETF"
    mems = members(period)
    rows = [[m.member_no, m.name, m.nic, m.earnings, m.etf] for m in mems]
    _sheet(ws, f"{c.name} – ETF contributions {period.label}",
           f"Employer no. {c.epf_zone}/{c.epf_employer_no or '—'} (ETF uses the EPF number) · {len(mems)} members",
           ["Member No", "Name", "NIC Number", "Total Earnings", "ETF 3%"], rows, {4, 5},
           c.theme()["brand"].lstrip("#"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def etf_half_year_excel(company, year, half):
    """ETF Form II style half-yearly schedule: each member's monthly 3% contributions."""
    months = range(1, 7) if half == 1 else range(7, 13)
    periods = {p.month: p for p in PayrollPeriod.objects.filter(company=company, year=year, month__in=months, status__in=["approved", "locked"])}
    data = {}
    for mth, p in periods.items():
        for m in members(p):
            key = m.slip.employee_id
            d = data.setdefault(key, {"m": m, "earn": Decimal(0), "etf": {}})
            d["etf"][mth] = m.etf
            d["earn"] += m.earnings
    rows = []
    for d in sorted(data.values(), key=lambda d: (d["m"].member_no or "").zfill(6)):
        m = d["m"]
        monthly = [d["etf"].get(x, Decimal(0)) for x in months]
        rows.append([m.member_no, m.name, m.nic, *monthly, sum(monthly, Decimal(0))])
    headers = ["Member No", "Name", "NIC Number"] + [calendar.month_abbr[x] for x in months] + ["Total ETF"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"ETF {year} H{half}"
    label = f"{'January–June' if half == 1 else 'July–December'} {year}"
    missing = [calendar.month_abbr[x] for x in months if x not in periods]
    _sheet(ws, f"{company.name} – ETF half-yearly return (Form II) {label}",
           f"Employer no. {company.epf_zone}/{company.epf_employer_no or '—'}"
           + (f" · No payroll for: {', '.join(missing)}" if missing else ""),
           headers, rows, set(range(4, 4 + len(months) + 1)), company.theme()["brand"].lstrip("#"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF schedules
def _pdf(period, title, subtitle, headers, rows, totals_row, widths, summary):
    c = period.frozen_company
    brand = colors.HexColor(c.theme()["brand"])
    soft = colors.HexColor(c.theme()["brand_soft"])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=title, author="Jay's HR")
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14, leading=18)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#555555"))
    head = ParagraphStyle("hd", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.white, alignment=1)
    story = [Paragraph(escape(title), h1), Paragraph(escape(subtitle), sub), Spacer(1, 4 * mm)]
    data = [[Paragraph(h, head) for h in headers], *rows, totals_row]
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand), ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f6f7f6")]),
        ("LINEBELOW", (0, 0), (-1, -2), 0.25, colors.HexColor("#dddddd")),
        ("BACKGROUND", (0, -1), (-1, -1), soft), ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ] + [("ALIGN", (i, 1), (i, -1), "RIGHT") for i, h in enumerate(headers) if h.startswith(("Total", "Employer", "Member's", "ETF", "Earn"))]))
    story += [t, Spacer(1, 5 * mm)]
    s = Table(summary, colWidths=[60 * mm, 40 * mm], hAlign="LEFT")
    s.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (1, 0), (1, -1), "Helvetica-Bold", 9),
                           ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("LINEABOVE", (0, 0), (-1, 0), 0.6, brand)]))
    story += [s, Spacer(1, 12 * mm),
              Paragraph("Certified correct: ..........................................    "
                        "Designation: ..............................    Date: ....................", sub)]

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(14 * mm, 8 * mm, f"{c.name} · {period.label} · Prepared by Jay's HR – PD Accounting Services (Pvt) Ltd")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _m(v):
    return f"{Decimal(v or 0):,.2f}"


def form_c_pdf(period, pay_mode="4", pay_ref="", pay_date=None):
    c = period.frozen_company
    mems = members(period)
    rows = [[m.member_no or "—", m.name, m.nic or "—", _m(m.earnings), _m(m.employer), _m(m.employee), _m(m.total)]
            for m in mems]
    tot = {k: sum((getattr(m, k) for m in mems), Decimal(0)) for k in ("earnings", "employer", "employee", "total")}
    mode = dict(PAY_MODES).get(str(pay_mode), "")
    return _pdf(
        period, f"EPF monthly contributions (Form C details) – {period.label}",
        f"<b>{c.name}</b> · {c.address.replace(chr(10), ', ')}<br/>EPF employer no. {c.epf_zone}/{c.epf_employer_no or '—'}"
        f" · Contribution period {period_code(period)} · {len(mems)} members",
        ["Member no.", "Name", "NIC no.", "Total earnings", "Employer 12%", "Member's 8%", "Total contribution"],
        rows, ["", "TOTAL", "", _m(tot["earnings"]), _m(tot["employer"]), _m(tot["employee"]), _m(tot["total"])],
        [22 * mm, 70 * mm, 38 * mm, 33 * mm, 33 * mm, 33 * mm, 36 * mm],
        [["Total contribution payable", f"Rs. {_m(tot['total'])}"], ["Payment method", mode or "—"],
         ["Payment reference", pay_ref or "—"], ["Payment date", (pay_date or date.today()).strftime("%d %b %Y")],
         ["Due by", "Last working day of the following month"]],
    )


def etf_pdf(period):
    c = period.frozen_company
    mems = members(period)
    rows = [[m.member_no or "—", m.name, m.nic or "—", _m(m.earnings), _m(m.etf)] for m in mems]
    te = sum((m.earnings for m in mems), Decimal(0))
    tt = sum((m.etf for m in mems), Decimal(0))
    form = "R1" if c.employees.filter(status="active").count() >= 15 else "R4"
    return _pdf(
        period, f"ETF contributions – {period.label}",
        f"<b>{c.name}</b> · Employer no. {c.epf_zone}/{c.epf_employer_no or '—'} (ETF uses the EPF number)"
        f" · Remittance form: {form}{' – online payment and e-return required' if form == 'R1' else ''}",
        ["Member no.", "Name", "NIC no.", "Total earnings", "ETF 3%"], rows,
        ["", "TOTAL", "", _m(te), _m(tt)], [25 * mm, 90 * mm, 45 * mm, 45 * mm, 40 * mm],
        [["Total ETF payable", f"Rs. {_m(tt)}"], ["Members", str(len(mems))],
         ["Due by", "Last working day of the following month"]],
    )


def filename(period, what, ext):
    return f"{what} - {period.frozen_company.name} - {period.label}.{ext}"


def stamp():
    return datetime.now().strftime("%d %b %Y %H:%M")
