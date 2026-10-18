from workforce.testing import company as test_company
from decimal import Decimal as D

from django.test import TestCase

from core.models import Company, Employee
from .engine import PayInputs, calculate, monthly_apit
from .models import PayrollInput, PayrollPeriod, run_payroll


class EngineTests(TestCase):
    def test_epf_etf_on_basic_only(self):
        r = calculate(PayInputs(basic_salary=D(30000), fixed_allowance=D(40000), bonus=D(10000)))
        self.assertEqual(r.epf_base, D("30000.00"))
        self.assertEqual(r.epf_employee, D("2400.00"))
        self.assertEqual(r.epf_employer, D("3600.00"))
        self.assertEqual(r.etf_employer, D("900.00"))
        self.assertEqual(r.gross_salary, D("80000.00"))
        self.assertEqual(r.net_salary, D("77600.00"))

    def test_epf_on_total_earnings(self):
        # basic 30,000 + COLA 5,000 (EPF) + travel 10,000 (not EPF) + commission 2,000 (EPF) + OT + bonus (not EPF)
        r = calculate(PayInputs(basic_salary=D(30000), epf_allowance=D(5000), fixed_allowance=D(10000),
                                epf_earnings=D(2000), ot_hours=D(8), bonus=D(5000)))
        self.assertEqual(r.epf_base, D("37000.00"))
        self.assertEqual(r.epf_employee, D("2960.00"))
        self.assertEqual(r.epf_employer, D("4440.00"))
        self.assertEqual(r.etf_employer, D("1110.00"))
        self.assertEqual(r.gross_salary, D("53500.00"))  # 30000+5000+10000+2000+1500 OT+5000

    def test_no_pay_reduces_epf_allowance_too(self):
        r = calculate(PayInputs(basic_salary=D(30000), epf_allowance=D(6000), fixed_allowance=D(9000),
                                no_pay_days=D(3)))
        self.assertEqual(r.no_pay_deduction, D("4500.00"))     # 45000/30*3
        self.assertEqual(r.epf_base, D("32400.00"))           # (30000+6000) - 3600

    def test_no_pay_reduces_epf_base(self):
        r = calculate(PayInputs(basic_salary=D(30000), fixed_allowance=D(15000), no_pay_days=D(3)))
        self.assertEqual(r.no_pay_deduction, D("4500.00"))  # (30000+15000)/30*3
        self.assertEqual(r.epf_base, D("27000.00"))
        self.assertEqual(r.epf_employee, D("2160.00"))

    def test_not_epf_eligible(self):
        r = calculate(PayInputs(basic_salary=D(40000), epf_eligible=False))
        self.assertEqual(r.epf_employee, 0)
        self.assertEqual(r.net_salary, D("40000.00"))

    def test_overtime(self):
        r = calculate(PayInputs(basic_salary=D(48000), ot_hours=D(10)))
        self.assertEqual(r.ot_amount, D("3000.00"))  # 48000/240*1.5*10

    def test_apit_bands(self):
        self.assertEqual(monthly_apit(D(150000)), D("0.00"))
        self.assertEqual(monthly_apit(D(200000)), D("3000.00"))  # 50,000 * 6%
        # 250,000: 83,333.33*6% + 16,666.67*18%
        self.assertEqual(monthly_apit(D(250000)), D("8000.00"))

    def test_deductions(self):
        r = calculate(PayInputs(basic_salary=D(40000), salary_advance=D(5000), welfare=D(500)))
        self.assertEqual(r.total_deductions, D("8700.00"))
        self.assertEqual(r.net_salary, D("31300.00"))


class RunPayrollTests(TestCase):
    def test_run(self):
        c = test_company(name="Test Co", code="TC")
        e = Employee.objects.create(company=c, emp_no="1", full_name="A", basic_salary=50000)
        Employee.objects.create(company=c, emp_no="2", full_name="B", basic_salary=1, status=Employee.RESIGNED)
        p = PayrollPeriod.objects.create(company=c, year=2026, month=10)
        PayrollInput.objects.create(period=p, employee=e, welfare=500)
        slips = run_payroll(p)
        self.assertEqual(len(slips), 1)
        self.assertEqual(slips[0].net_salary, D("45500.00"))


import io  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from core.models import User  # noqa: E402
from .models import start_period  # noqa: E402


class CarryForwardAndExportTests(TestCase):
    def setUp(self):
        self.c = test_company(name="Test Co", code="TC")
        self.e1 = Employee.objects.create(company=self.c, emp_no="1", full_name="Anu", basic_salary=30000,
                                          fixed_allowance=20000)
        self.e2 = Employee.objects.create(company=self.c, emp_no="2", full_name="Bala", basic_salary=30000)
        self.hr = User.objects.create_user("hr", password="pw", role=User.CLIENT_HR, company=self.c,
                                           must_change_password=False)
        self.sep, _ = start_period(self.c, 2026, 9)
        i = self.sep.inputs.get(employee=self.e1)
        i.welfare, i.loan_recovery, i.remarks = 500, 2000, "Loan 3/10"
        i.save()

    def test_next_month_starts_with_last_months_inputs(self):
        self.client.force_login(self.hr)
        r = self.client.post(f"/payroll/company/{self.c.pk}/", {"year": 2026, "month": 10, "carry_forward": "on"})
        oct_ = PayrollPeriod.objects.get(month=10)
        self.assertEqual(oct_.copied_from, self.sep)
        i = oct_.inputs.get(employee=self.e1)
        self.assertEqual((i.welfare, i.loan_recovery, i.remarks), (500, 0, ""))
        self.assertContains(self.client.get(r.url), "carried forward from")

    def test_new_employee_and_opt_out(self):
        e3 = Employee.objects.create(company=self.c, emp_no="3", full_name="Chan", basic_salary=30000)
        oct_, _ = start_period(self.c, 2026, 10)
        self.assertEqual(oct_.inputs.get(employee=e3).welfare, 0)          # nothing to copy for a new joiner
        nov, _ = start_period(self.c, 2026, 11, carry_forward=False)
        self.assertEqual(nov.inputs.get(employee=self.e1).welfare, 0)
        self.assertIsNone(nov.copied_from)

    def test_copy_previous_again(self):
        oct_, _ = start_period(self.c, 2026, 10)
        i = oct_.inputs.get(employee=self.e1)
        i.welfare = 0
        i.save()
        self.client.force_login(self.hr)
        self.client.post(f"/payroll/{oct_.pk}/copy_previous/")
        self.assertEqual(oct_.inputs.get(employee=self.e1).welfare, 500)

    def test_excel_export(self):
        run_payroll(self.sep)
        self.client.force_login(self.hr)
        r = self.client.get(f"/payroll/{self.sep.pk}/export.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        ws = load_workbook(io.BytesIO(r.content)).active
        values = [[c.value for c in row] for row in ws.iter_rows()]
        flat = [v for row in values for v in row if v is not None]
        self.assertIn("Anu", flat)
        self.assertIn(45100.0, flat)    # Anu net: 50000 - 2400 EPF - 500 welfare - 2000 loan
        self.assertTrue(any(isinstance(v, str) and v.startswith("=SUM(") for v in flat))

    def test_pdf_export_and_draft(self):
        self.client.force_login(self.hr)
        r = self.client.get(f"/payroll/{self.sep.pk}/export.pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_export_permissions(self):
        other = test_company(name="Other", code="OT")
        p2, _ = start_period(other, 2026, 9)
        emp_user = User.objects.create_user("e", password="pw", role=User.EMPLOYEE, company=self.c,
                                            employee=self.e1, must_change_password=False)
        self.client.force_login(self.hr)
        self.assertEqual(self.client.get(f"/payroll/{p2.pk}/export.xlsx").status_code, 403)
        self.client.force_login(emp_user)
        self.assertEqual(self.client.get(f"/payroll/{self.sep.pk}/export.pdf").status_code, 403)


import zipfile  # noqa: E402
from datetime import date as _date  # noqa: E402

from . import statutory  # noqa: E402


class StatutoryTests(TestCase):
    def setUp(self):
        self.c = test_company(name="Test Co", code="TC", epf_zone="A", epf_employer_no="12345",
                                        epf_district_office="CO")
        self.e1 = Employee.objects.create(company=self.c, emp_no="1", full_name="Kamal Perera", nic="199012345678",
                                          epf_no="7", basic_salary=30000, epf_allowance=5000, fixed_allowance=10000,
                                          join_date=_date(2026, 10, 1))
        self.e2 = Employee.objects.create(company=self.c, emp_no="2", full_name="Nimal Silva", nic="851234567V",
                                          epf_no="3", basic_salary=40000, join_date=_date(2025, 1, 1))
        Employee.objects.create(company=self.c, emp_no="3", full_name="No EPF", basic_salary=40000, epf_eligible=False)
        self.p = PayrollPeriod.objects.create(company=self.c, year=2026, month=10)
        run_payroll(self.p)
        self.p.refresh_from_db()
        self.p.status = "approved"
        self.p.save()

    def test_members_and_status(self):
        mems = statutory.members(self.p)
        self.assertEqual([m.member_no for m in mems], ["3", "7"])
        kamal = mems[1]
        self.assertEqual((kamal.surname, kamal.initials, kamal.status), ("PERERA", "K", "N"))
        self.assertEqual(kamal.earnings, D("35000.00"))          # basic + EPF-able allowance, not the 10,000
        self.assertEqual(kamal.total, D("7000.00"))              # 20%

    def test_evemc_fixed_width(self):
        lines = statutory.evemc(self.p).rstrip("\r\n").split("\r\n")
        self.assertEqual(len(lines), 2)
        line = lines[1]
        self.assertEqual(len(line), 154)
        self.assertEqual(line[0:20].strip(), "199012345678")
        self.assertEqual(line[60:80].strip(), "K")
        self.assertEqual(line[80:86], "000007")
        self.assertEqual(line[86:96], "0007000.00")      # total contribution
        self.assertEqual(line[96:106], "0004200.00")     # employer 12%
        self.assertEqual(line[106:116], "0002800.00")    # member 8%
        self.assertEqual(line[116:128], "000035000.00")  # total earnings
        self.assertEqual(line[128:142], "NA012345202610")

    def test_evemp(self):
        line = statutory.evemp(self.p, pay_mode="4", pay_ref="7010-123", pay_date=_date(2026, 11, 20)).strip()
        self.assertEqual(len(line), 63)
        self.assertTrue(line.startswith("A01234520261001000015000.00000024"))   # 7000 + 8000 = 15000
        self.assertTrue(line.endswith("20261120CO"))

    def test_zip_and_downloads(self):
        hr = User.objects.create_user("hr", password="pw", role=User.CLIENT_HR, company=self.c,
                                      must_change_password=False)
        self.client.force_login(hr)
        r = self.client.get(f"/payroll/{self.p.pk}/epf-etf/epf-return/?pay_mode=1&pay_ref=CHQ55&submission=1")
        names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        self.assertIn("EVEMC.TXT", names)
        for what in ("form-c", "etf-pdf"):
            self.assertTrue(self.client.get(f"/payroll/{self.p.pk}/epf-etf/{what}/").content.startswith(b"%PDF"))
        for what in ("epf-excel", "etf-excel"):
            self.assertEqual(self.client.get(f"/payroll/{self.p.pk}/epf-etf/{what}/").status_code, 200)
        page = self.client.get(f"/payroll/{self.p.pk}/epf-etf/")
        self.assertContains(page, "15,000.00")
        r = self.client.get(f"/payroll/company/{self.c.pk}/etf-half-year/?year=2026&half=2")
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertIn("Oct", [c.value for c in ws[4]])

    def test_problems_listed(self):
        self.c.epf_employer_no = ""
        self.c.save()
        self.e2.nic = ""
        self.e2.save()
        issues = statutory.problems(self.p)
        self.assertFalse(any("employer number" in i for i in issues))  # company registration is frozen
        self.assertFalse(any("Nimal Silva" in i for i in issues))  # historical NIC is frozen
