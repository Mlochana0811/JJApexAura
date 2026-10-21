"""Salary sheet downloads (Excel and PDF) for one payroll month."""

import io
from html import escape
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from core.models import Employee

from .engine import calculate
from workforce.reports import safe_cell
from .models import pay_inputs

# (key, heading, is_money)
COLUMNS = [
    ("emp_no", "Emp. no.", False), ("name", "Employee", False), ("designation", "Designation", False),
    ("epf_no", "EPF no.", False),
    ("basic_salary", "Basic salary", True), ("epf_allowance", "EPF-able allowances", True),
    ("fixed_allowance", "Other allowances", True),
    ("no_pay_deduction", "No-pay", True), ("ot_amount", "Overtime", True), ("bonus", "Bonus", True),
    ("epf_earnings", "EPF-able earnings", True), ("other_earnings", "Other earnings", True),
    ("gross_salary", "Gross salary", True), ("epf_base", "EPF/ETF base", True),
    ("epf_employee", "EPF 8%", True), ("apit", "APIT", True), ("salary_advance", "Salary advance", True),
    ("welfare", "Welfare", True), ("loan_recovery", "Loan", True), ("other_deductions", "Other deductions", True),
    ("total_deductions", "Total deductions", True), ("net_salary", "Net salary", True),
    ("epf_employer", "EPF 12%", True), ("etf_employer", "ETF 3%", True), ("employer_cost", "Cost to company", True),
    ("bank", "Bank", False), ("account", "Account no.", False),
]
MONEY_KEYS = [k for k, _, m in COLUMNS if m]


def salary_rows(period):
    """Rows for the sheet. Uses processed payslips if there are any; otherwise a draft from current inputs."""
    slips = list(period.payslips.select_related("employee").order_by("employee__emp_no"))
    company = period.frozen_company
    rows = []
    if slips and period.status in ("processed", "approved", "locked"):
        for s in slips:
            rows.append(_row(s.frozen_employee, {k: getattr(s, k) for k in MONEY_KEYS}))
        return rows, False
    for inp in period.inputs.select_related("employee").filter(employee__status=Employee.ACTIVE) \
            .order_by("employee__emp_no"):
        e = inp.employee
        r = calculate(pay_inputs(e, inp, company))
        rows.append(_row(e, {k: getattr(r, k) for k in MONEY_KEYS}))
    return rows, True


def _row(e, money):
    return {"emp_no": e.emp_no, "name": e.full_name, "designation": e.designation, "epf_no": e.epf_no,
            "bank": f"{e.bank_name} {e.bank_branch}".strip(), "account": e.bank_account, **money}


def totals(rows):
    return {k: sum((r[k] for r in rows), Decimal(0)) for k in MONEY_KEYS}


def filename(period, ext):
    return f"Salary sheet - {period.frozen_company.name} - {period.label}.{ext}"


# ---------------------------------------------------------------- Excel
def to_excel(period):
    slips=list(period.payslips.all())
    if period.status in ('processed','approved','locked') and slips and all(s.snapshot.get('salary_template',{}).get('config') for s in slips) and all(s.snapshot['salary_template']['config']==slips[0].snapshot['salary_template']['config'] for s in slips):
        return template_payroll_excel(period,slips)
    rows, draft = salary_rows(period)
    brand = period.frozen_company.theme()["brand"].lstrip("#")
    wb = Workbook()
    ws = wb.active
    ws.title = period.label[:31]
    ws["A1"] = period.frozen_company.name
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Salary sheet – {period.label}" + ("  (DRAFT – not yet processed)" if draft else "")
    ws["A2"].font = Font(size=11, bold=draft, color="B42318" if draft else "444444")
    ws["A3"] = f"Status: {period.get_status_display()} · Downloaded {datetime.now():%d %b %Y %H:%M} · Jay's HR"
    ws["A3"].font = Font(size=9, color="777777")

    head_row = 5
    thin = Side(style="thin", color="DDDDDD")
    for c, (_, title, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(head_row, c, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=brand)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[head_row].height = 30

    for r_i, row in enumerate(rows, head_row + 1):
        for c, (key, _, is_money) in enumerate(COLUMNS, 1):
            v = row[key]
            cell = ws.cell(r_i, c, float(v) if is_money else safe_cell(v))
            cell.border = Border(bottom=thin)
            if is_money:
                cell.number_format = "#,##0.00"

    last = head_row + len(rows)
    tot_row = last + 1
    ws.cell(tot_row, 2, "TOTAL").font = Font(bold=True)
    for c, (key, _, is_money) in enumerate(COLUMNS, 1):
        if is_money:
            col = get_column_letter(c)
            cell = ws.cell(tot_row, c, f"=SUM({col}{head_row + 1}:{col}{last})" if rows else 0)
            cell.font = Font(bold=True)
            cell.number_format = "#,##0.00"
            cell.fill = PatternFill("solid", fgColor="F1F3F2")
            cell.border = Border(top=Side(style="thin", color="888888"))

    # Statutory summary
    col = {k: get_column_letter(i) for i, (k, _, _) in enumerate(COLUMNS, 1)}
    s = tot_row + 2
    ws.cell(s, 2, "Payments due").font = Font(bold=True)
    summary = [
        ("Net salaries to bank", f"={col['net_salary']}{tot_row}"),
        ("EPF payable (8% + 12%)", f"={col['epf_employee']}{tot_row}+{col['epf_employer']}{tot_row}"),
        ("ETF payable (3%)", f"={col['etf_employer']}{tot_row}"),
        ("APIT payable", f"={col['apit']}{tot_row}"),
    ]
    for i, (label, formula) in enumerate(summary, s + 1):
        ws.cell(i, 2, label)
        cell = ws.cell(i, 3, formula)
        cell.number_format = "#,##0.00"
        cell.font = Font(bold=True)
    ws.cell(s + 6, 2, "EPF and ETF are calculated on total earnings: basic + EPF-able allowances and earnings (not OT, bonus or other allowances).").font = Font(italic=True, size=9, color="777777")

    widths = {"emp_no": 9, "name": 28, "designation": 22, "epf_no": 8, "bank": 26, "account": 18}
    for c, (key, _, _) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(key, 13)
    ws.freeze_panes = ws.cell(head_row + 1, 3)
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
PDF_COLUMNS = [
    ("emp_no", "No."), ("name", "Employee"), ("basic_salary", "Basic"), ("fixed_allowance", "Allowances"),
    ("no_pay_deduction", "No-pay"), ("extra", "OT / bonus / other"), ("gross_salary", "Gross"),
    ("epf_base", "EPF base"),
    ("epf_employee", "EPF 8%"), ("apit", "APIT"), ("other_ded", "Advance / welfare / loan / other"),
    ("net_salary", "Net salary"), ("epf_employer", "EPF 12%"), ("etf_employer", "ETF 3%"),
]


def _m(v):
    v = Decimal(v or 0)
    return "–" if v == 0 else f"{v:,.2f}"


def to_pdf(period):
    rows, draft = salary_rows(period)
    t = totals(rows)
    company = period.frozen_company
    brand = colors.HexColor(company.theme()["brand"])
    soft = colors.HexColor(company.theme()["brand_soft"])

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm,
                            title=filename(period, "pdf")[:-4], author="Jay's HR")
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, leading=18)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=10, leading=13, textColor=colors.HexColor("#555555"))
    cellst = ParagraphStyle("cell", fontName="Helvetica", fontSize=7.5, leading=9)
    headst = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white,
                            alignment=1)

    story = []
    if company.logo:
        try:
            img = Image(company.logo.path)
            ratio = img.imageWidth / img.imageHeight
            img.drawHeight, img.drawWidth = 12 * mm, 12 * mm * ratio
            img.hAlign = "LEFT"
            story.append(img)
            story.append(Spacer(1, 3 * mm))
        except Exception:
            pass
    story.append(Paragraph(escape(company.name), h1))
    title = f"Salary sheet – {period.label}"
    if draft:
        title += ' <font color="#b42318"><b>DRAFT – not yet processed</b></font>'
    story.append(Paragraph(title, sub))
    story.append(Paragraph(f"Status: {period.get_status_display()} · Downloaded {datetime.now():%d %b %Y %H:%M}", sub))
    story.append(Spacer(1, 5 * mm))

    data = [[Paragraph(h, headst) for _, h in PDF_COLUMNS]]
    for r in rows:
        extra = r["ot_amount"] + r["bonus"] + r["other_earnings"] + r["epf_earnings"]
        other_ded = r["salary_advance"] + r["welfare"] + r["loan_recovery"] + r["other_deductions"]
        data.append([r["emp_no"], Paragraph(escape(r["name"]), cellst), _m(r["basic_salary"]),
                     _m(r["fixed_allowance"] + r["epf_allowance"]),
                     _m(r["no_pay_deduction"]), _m(extra), _m(r["gross_salary"]), _m(r["epf_base"]),
                     _m(r["epf_employee"]),
                     _m(r["apit"]), _m(other_ded), _m(r["net_salary"]), _m(r["epf_employer"]),
                     _m(r["etf_employer"])])
    data.append(["", "TOTAL", _m(t["basic_salary"]), _m(t["fixed_allowance"] + t["epf_allowance"]),
                 _m(t["no_pay_deduction"]), _m(t["ot_amount"] + t["bonus"] + t["other_earnings"] + t["epf_earnings"]),
                 _m(t["gross_salary"]), _m(t["epf_base"]), _m(t["epf_employee"]), _m(t["apit"]),
                 _m(t["salary_advance"] + t["welfare"] + t["loan_recovery"] + t["other_deductions"]),
                 _m(t["net_salary"]), _m(t["epf_employer"]), _m(t["etf_employer"])])

    widths = [13 * mm, 44 * mm] + [17 * mm, 18 * mm, 15 * mm, 19 * mm, 18 * mm, 17 * mm, 15 * mm, 14 * mm,
                                   23 * mm, 20 * mm, 15 * mm, 15 * mm]
    tbl = Table(data, colWidths=widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 7.5),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f6f7f6")]),
        ("LINEBELOW", (0, 0), (-1, -2), 0.25, colors.HexColor("#dddddd")),
        ("BACKGROUND", (0, -1), (-1, -1), soft),
        ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 7.5),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, brand),
        ("FONT", (11, 1), (11, -1), "Helvetica-Bold", 7.5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 6 * mm))

    summary = Table([
        ["Payments due", ""],
        ["Net salaries to bank", f"Rs. {t['net_salary']:,.2f}"],
        ["EPF payable (8% + 12%)", f"Rs. {t['epf_employee'] + t['epf_employer']:,.2f}"],
        ["ETF payable (3%)", f"Rs. {t['etf_employer']:,.2f}"],
        ["APIT payable", f"Rs. {t['apit']:,.2f}"],
        ["Total cost to company", f"Rs. {t['employer_cost']:,.2f}"],
    ], colWidths=[55 * mm, 35 * mm], hAlign="LEFT")
    summary.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
        ("TEXTCOLOR", (0, 0), (-1, 0), brand), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("FONT", (1, 1), (1, -1), "Helvetica-Bold", 8.5),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, brand), ("TOPPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(summary)
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("EPF and ETF are calculated on total earnings: basic + EPF-able allowances and earnings (not OT, bonus or other allowances).", sub))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(12 * mm, 8 * mm, f"{company.name} · {period.label} · Prepared by Jay's HR – PD Accounting Services (Pvt) Ltd")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def template_payroll_excel(period,slips):
    """Template-based processed report, with APIT and final net added explicitly."""
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from workforce.template_payroll import FIELDS, workbook_bytes
    config=slips[0].snapshot['salary_template']['config']
    if any(s.snapshot['salary_template']['config']!=config for s in slips):
        raise ValueError('Payroll contains mixed template configurations')
    rows=[]
    for slip in slips:
        t=slip.snapshot['salary_template'];e=slip.frozen_employee
        row=dict.fromkeys(FIELDS,'0');row.update(t['monthly_inputs']);row.update(basic_salary=t['basic_salary'],name=e.full_name,emp_no=e.emp_no)
        for i,a in enumerate(t['allowances']):row[f'allowance_{i}']=a['amount']
        rows.append(row)
    proxy=SimpleNamespace(company=period.frozen_company,company_name=period.frozen_company.name,year=period.year,month=period.month,status=period.status,template_config={**config,"custom_columns":[]},rows=rows)
    w=load_workbook(io.BytesIO(workbook_bytes(proxy)));ws=w['Payroll']
    ws['A5']='Processed payroll snapshot. Contributions reflect the reviewed payroll policy and eligibility. X: APIT; Y: final net pay.'
    from copy import copy
    for col in (24,25):
        ws.cell(8,col)._style=copy(ws.cell(8,23)._style)
        for row in range(9,9+len(slips)):ws.cell(row,col)._style=copy(ws.cell(row,23)._style)
    ws['T8']='EPF employee';ws['U8']='EPF employer';ws['V8']='ETF employer';ws['X8']='APIT';ws['Y8']='Net pay after APIT'
    for i,s in enumerate(slips,9):
        for col,value in enumerate([s.gross_salary,s.epf_base,s.epf_employee,s.epf_employer,s.etf_employer,s.net_salary+s.apit,s.apit,s.net_salary],18):
            ws.cell(i,col,value).number_format='#,##0.00'
        ws.cell(i,15,s.loan_recovery)
    end=8+len(slips)
    for col in range(18,26):
        letter=get_column_letter(col);ws.cell(4,col,f'=SUM({letter}9:{letter}{end})')
    ws.column_dimensions['X'].width=18;ws.column_dimensions['Y'].width=23
    ws.print_area=f'A1:Y{end}';ws.auto_filter.ref=f'A8:Y{end}'
    stream=io.BytesIO();w.save(stream);return stream.getvalue()
