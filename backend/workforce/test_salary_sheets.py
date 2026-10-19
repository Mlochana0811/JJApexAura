import io, json
from django.test import TestCase
from openpyxl import load_workbook
from core.models import Company, Employee, User
from .models import SalarySheet

class SalarySheetTests(TestCase):
    def setUp(self):
        self.c=Company.objects.create(name='Example',code='E');self.other=Company.objects.create(name='Other',code='O')
        self.e=Employee.objects.create(company=self.c,emp_no='1',full_name='Example person',basic_salary=70000,bank_account='001234')
        self.u=User.objects.create_user('hr',role='client_hr',company=self.c,must_change_password=False);self.client.force_login(self.u)
    def data(self):return {'company':self.c.pk,'year':2026,'month':10,'rows':[{'employee_id':self.e.pk,'basic_salary':'70000','epf_salary':'30000','epf_rate':'8','leave_deductions':'1000','welfare':'500','salary_advance':'5000','format_percent':'20'}]}
    def post(self,path,data):return self.client.post(path,json.dumps(data),content_type='application/json')
    def create(self):
        r=self.post('/api/salary-sheets/',self.data());self.assertEqual(r.status_code,201,r.content);return r.json()['item']
    def test_calculation(self):
        d=self.data();d['rows'][0].update(epf='0',salary_amount='999999');r=self.post('/api/salary-sheets/',d)
        self.assertEqual(r.status_code,201,r.content);row=r.json()['item']['rows'][0];self.assertEqual(row['epf'],'2400.00');self.assertEqual(row['salary_amount'],'61100.00')
    def test_permissions(self):
        s=self.create();other=User.objects.create_user('other',role='client_hr',company=self.other,must_change_password=False);self.client.force_login(other)
        for path in (f'/api/salary-sheets/{s["id"]}/',f'/api/salary-sheets/{s["id"]}/export/'):
            self.assertEqual(self.client.get(path).status_code,404)
        self.assertEqual(self.post('/api/salary-sheets/',self.data()).status_code,404)
        d=self.data();d['company']=self.other.pk;self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
        for role in ('manager','employee'):
            u=User.objects.create_user(role,role=role,company=self.c,must_change_password=False);self.client.force_login(u)
            for path in ('/api/salary-sheets/','/api/payroll-settings/',f'/api/schema/payroll-settings/{self.e.pk}/',f'/api/salary-sheets/{s["id"]}/export/'):
                self.assertEqual(self.client.get(path).status_code,403)
    def test_validation(self):
        for k,v in [('basic_salary','NaN'),('epf_salary','Infinity'),('welfare','-1'),('epf_rate','101')]:
            d=self.data();d['rows'][0][k]=v;self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
        d=self.data();d['rows']*=2;self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
        d=self.data();d['month']=13;self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
    def test_locking_and_revision(self):
        s=self.create();path=f'/api/salary-sheets/{s["id"]}/';d=self.data();d['revision']=0
        self.assertEqual(self.post(path,d).status_code,409);d['revision']=1;d['finalize']=True
        self.assertEqual(self.post(path,d).status_code,200);d['revision']=2;self.assertEqual(self.post(path,d).status_code,400)
        self.assertEqual(SalarySheet.objects.get(pk=s['id']).status,'finalized')
    def test_negative_salary_draft_only(self):
        d=self.data();d['rows'][0]['salary_advance']='90000';d['finalize']=True
        self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400);d['finalize']=False
        self.assertEqual(self.post('/api/salary-sheets/',d).status_code,201)
    def test_export_and_snapshot(self):
        s=self.create();self.e.full_name='Changed';self.e.save()
        r=self.client.get(f'/api/salary-sheets/{s["id"]}/export/');self.assertEqual(r.status_code,200)
        ws=load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([ws.cell(4,c).value for c in range(2,10)],['Name','Basic Salary','EPF Salary','Leave Deductions','Welfare','Salary Advance','EPF','Salary Amount'])
        self.assertEqual(ws['B5'].value,'Example person');self.assertEqual(ws['H5'].value,'=ROUND(D5*8.00/100,2)');self.assertEqual(ws['I5'].value,'=ROUND(C5-E5-F5-G5-H5,2)');self.assertEqual(ws['I7'].value,'=SUM(I5:I5)');self.assertEqual(ws['M5'].value,.2)
    def test_formula_injection(self):
        self.e.full_name='=1+2';self.e.save();s=self.create();r=self.client.get(f'/api/salary-sheets/{s["id"]}/export/')
        ws=load_workbook(io.BytesIO(r.content)).active;self.assertEqual(ws['B5'].data_type,'s')
    def test_fields_moved_without_data_loss(self):
        fields=self.client.get('/api/schema/employees/').json()['fields'];self.assertNotIn('bank_account',[f['name'] for f in fields]);self.assertNotIn('basic_salary',[f['name'] for f in fields])
        fields=self.client.get(f'/api/schema/payroll-settings/{self.e.pk}/').json()['fields'];self.assertIn('bank_account',[f['name'] for f in fields]);d={f['name']:f['initial'] for f in fields};d['bank_name']='Example bank'
        r=self.post(f'/api/payroll-settings/{self.e.pk}/',d);self.assertEqual(r.status_code,200,r.content);self.e.refresh_from_db();self.assertEqual(self.e.bank_account,'001234');self.assertEqual(self.e.basic_salary,70000)
    def test_duplicate_month(self):
        self.create();self.assertEqual(self.post('/api/salary-sheets/',self.data()).status_code,400);self.assertEqual(SalarySheet.objects.count(),1)
