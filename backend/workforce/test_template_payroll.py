import io,json
from decimal import Decimal as D
from pathlib import Path
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from core.models import Employee,User
from payroll.models import PayrollInput,run_payroll,PayrollPeriod
from .testing import company
from .template_payroll import default_config,totals,FIELDS

class TemplatePayrollTests(TestCase):
 def setUp(self):
  self.c=company(name='Example',code='EX');self.e=Employee.objects.create(company=self.c,emp_no='001',full_name='Example',basic_salary=30000,apit_applicable=False)
  self.u=User.objects.create_user('hr',role='client_hr',company=self.c,must_change_password=False);self.client.force_login(self.u)
 def post(self,url,data):return self.client.post(url,json.dumps(data),content_type='application/json')
 def data(self):
  row=dict.fromkeys(FIELDS,'0');row.update(employee_id=self.e.pk,basic_salary='30000',allowance_0='10000',allowance_3='5000',no_pay_days='2',ot_hours='4',bonus='1000',welfare='500',remarks='Example')
  return {'company':self.c.pk,'year':2026,'month':10,'template_config':default_config(self.c),'rows':[row]}
 def test_formula_and_flags(self):
  d=self.data();r=totals(d['rows'][0],d['template_config']);self.assertEqual(r['gross_salary'],'43750.00');self.assertEqual(r['epf_base'],'37333.33');self.assertEqual(r['salary_amount'],'40263.33')
  d['template_config']['allowances'][3]['epf']=True
  self.assertEqual(totals(d['rows'][0],d['template_config'])['epf_base'],'42000.00')
 def test_save_export_apply_and_process(self):
  d=self.data();d['finalize']=True;r=self.post('/api/salary-sheets/',d);self.assertEqual(r.status_code,201,r.content);s=r.json()['item']
  export=self.client.get(f'/api/salary-sheets/{s["id"]}/export/');self.assertEqual(export.status_code,200)
  w=load_workbook(io.BytesIO(export.content));self.assertEqual(w['Payroll']['A9'].value,'001');self.assertEqual(w['Payroll']['D9'].value,10000);self.assertIn('SUMPRODUCT',w['Payroll']['S9'].value)
  r=self.post(f'/api/salary-sheets/{s["id"]}/apply/',{});self.assertEqual(r.status_code,200,r.content)
  p=PayrollPeriod.objects.get(pk=r.json()['period_id']);run_payroll(p);slip=p.payslips.get();self.assertEqual(slip.gross_salary,D('43750'));self.assertEqual(slip.net_salary,D(s['rows'][0]['salary_amount']));self.assertEqual(slip.snapshot['salary_template']['allowances'][0]['name'],'Cost of living allowance')
  from payroll.exports import to_excel
  p.status='processed';p.save()
  report=load_workbook(io.BytesIO(to_excel(p)));self.assertEqual(report['Payroll']['X8'].value,'APIT');self.assertEqual(D(str(report['Payroll']['Y9'].value)),slip.net_salary)
  p.status='approved';p.save();self.assertEqual(self.post(f'/api/salary-sheets/{s["id"]}/apply/',{}).status_code,400)
 def test_upload_validates_identity_and_ignores_check_formulas(self):
  w=load_workbook(Path(__file__).parent/'templates_data'/'salary-upload.xlsx');ws=w['Payroll'];ws['A9']='001';ws['C9']=30000;ws['D9']=5000
  def upload():
   b=io.BytesIO();w.save(b);return self.client.post('/api/salary-sheets/upload/',{'company':self.c.pk,'year':2026,'month':10,'file':SimpleUploadedFile('salary.xlsx',b.getvalue())})
  r=upload();self.assertEqual(r.status_code,200,r.content);self.assertEqual(r.json()['rows'][0]['epf_base'],'35000.00');self.assertFalse(PayrollInput.objects.exists())
  ws['B2']='WRONG';self.assertEqual(upload().status_code,400);ws['B2']=None
  ws['A10']='001';ws['C10']=30000;self.assertEqual(upload().status_code,400);ws['A10']=None;ws['C10']=None
  ws['C9']='=100+200';self.assertEqual(upload().status_code,400)
 def test_isolation_and_invalid_inputs(self):
  d=self.data();d['rows'][0]['no_pay_days']='31';self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
  d=self.data();d['template_config']['allowances'][0]['epf']='yes';self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
  other=company(name='Other',code='OT');d=self.data();d['company']=other.pk;self.assertEqual(self.post('/api/salary-sheets/',d).status_code,404)

 def test_custom_columns_visibility_and_export(self):
  d=self.data();cfg=d['template_config'];cfg['custom_columns']=[{'key':'custom_note','name':'Department note'}];cfg['column_visibility']={'basic_salary':False,'custom_note':True}
  d['rows'][0]['custom_values']={'custom_note':'Reviewed'}
  r=self.post('/api/salary-sheets/',d);self.assertEqual(r.status_code,201,r.content);sheet=r.json()['item']
  self.assertEqual(sheet['rows'][0]['custom_values'],{'custom_note':'Reviewed'})
  self.assertEqual(sheet['rows'][0]['salary_amount'],'40263.33')
  export=self.client.get(f'/api/salary-sheets/{sheet["id"]}/export/')
  ws=load_workbook(io.BytesIO(export.content))['Payroll']
  self.assertTrue(ws.column_dimensions['C'].hidden);self.assertEqual(ws['X8'].value,'Department note');self.assertEqual(ws['X9'].value,'Reviewed')
 def test_invalid_column_config(self):
  d=self.data();d['template_config']['column_visibility']={'basic_salary':'No'}
  self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
  d=self.data();d['template_config']['custom_columns']=[{'key':'basic_salary','name':'Override'}]
  self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)

 def test_removed_columns_clear_values_and_preserve_headers(self):
  d=self.data();d['template_config']['allowances'][0]['enabled']=False
  d['template_config']['removed_columns']=['bonus','remarks']
  d['template_config']['header_names']={'basic_salary':'Monthly basic'}
  r=self.post('/api/salary-sheets/',d);self.assertEqual(r.status_code,201,r.content)
  sheet=r.json()['item'];self.assertEqual(sheet['rows'][0]['allowance_0'],'0.00');self.assertEqual(sheet['rows'][0]['bonus'],'0.00');self.assertEqual(sheet['rows'][0]['remarks'],'')
  self.assertEqual(sheet['template_config']['header_names']['basic_salary'],'Monthly basic')
  self.assertFalse(sheet['template_config']['allowances'][0]['enabled'])
  export=self.client.get(f'/api/salary-sheets/{sheet["id"]}/export/')
  ws=load_workbook(io.BytesIO(export.content))['Payroll'];self.assertEqual(ws['C8'].value,'Monthly basic');self.assertTrue(ws.column_dimensions['D'].hidden)
  response=self.client.post('/api/salary-sheets/upload/',{'company':self.c.pk,'year':2026,'month':10,'file':SimpleUploadedFile('salary.xlsx',export.content)})
  self.assertEqual(response.status_code,200,response.content)
  self.assertEqual(response.json()['template_config']['header_names']['basic_salary'],'Monthly basic')

 def test_separate_etf_flags_and_payroll(self):
  d=self.data();d['template_config']['allowances'][3]['etf']=True;d['finalize']=True
  response=self.post('/api/salary-sheets/',d);self.assertEqual(response.status_code,201,response.content);sheet=response.json()['item']
  self.assertEqual(sheet['rows'][0]['epf_base'],'37333.33');self.assertEqual(sheet['rows'][0]['etf_employer'],'1260.00')
  exported=self.client.get(f'/api/salary-sheets/{sheet["id"]}/export/')
  ws=load_workbook(io.BytesIO(exported.content))['Payroll'];self.assertEqual(ws['G6'].value,'Yes');self.assertEqual(ws['G7'].value,'No');self.assertIn('$D$6',ws['V9'].value)
  uploaded=self.client.post('/api/salary-sheets/upload/',{'company':self.c.pk,'year':2026,'month':10,'file':SimpleUploadedFile('salary.xlsx',exported.content)})
  self.assertEqual(uploaded.status_code,200,uploaded.content);self.assertTrue(uploaded.json()['template_config']['allowances'][3]['etf'])
  r=self.post(f'/api/salary-sheets/{sheet["id"]}/apply/',{});self.assertEqual(r.status_code,200,r.content)
  period=PayrollPeriod.objects.get(pk=r.json()['period_id']);run_payroll(period)
  self.assertEqual(period.payslips.get().etf_employer,D('1260'))
 def test_archive_revision_restore_and_isolation(self):
  d=self.data();d['finalize']=True;r=self.post('/api/salary-sheets/',d);sheet=r.json()['item'];url=f'/api/salary-sheets/{sheet["id"]}/'
  def patch(revision,archived):return self.client.patch(url,json.dumps({'revision':revision,'archived':archived}),content_type='application/json')
  self.assertEqual(patch(0,True).status_code,409)
  r=patch(sheet['revision'],True);self.assertEqual(r.status_code,200,r.content)
  self.assertEqual(self.client.get('/api/salary-sheets/').json()['count'],0)
  self.assertEqual(self.client.get('/api/salary-sheets/?archived=true').json()['count'],1)
  self.assertEqual(self.post(f'{url}apply/',{}).status_code,400)
  self.assertEqual(patch(r.json()['item']['revision'],False).status_code,200)
  self.assertEqual(self.client.get('/api/salary-sheets/').json()['count'],1)
  other=company(name='Other',code='ZZ');u=User.objects.create_user('otherhr',role='client_hr',company=other,must_change_password=False);self.client.force_login(u)
  self.assertEqual(patch(3,True).status_code,404)
 def test_food_exclusion_requires_review(self):
  d=self.data();d['finalize']=True;d['template_config']['allowances'][1]['epf']=False
  self.assertEqual(self.post('/api/salary-sheets/',d).status_code,400)
