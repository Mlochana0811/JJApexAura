"""Manual salary-sheet workflow matching the supplied B:I workbook layout."""
import calendar
import io
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from django import forms
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_http_methods
from core.models import Employee
from .api import endpoint, body, companies
from .models import SalarySheet
from .services import require_hr, audit, plain

MONEY = ('basic_salary','epf_salary','leave_deductions','welfare','salary_advance')
OPTIONAL = ('appraisal','new_salary','format_percent')

class SheetHeader(forms.Form):
    company = forms.IntegerField(min_value=1)
    year = forms.IntegerField(min_value=2000,max_value=2100)
    month = forms.IntegerField(min_value=1,max_value=12)

def amount(value, label, optional=False):
    if optional and value in (None,''):return None
    try:
        n=Decimal(str(value))
        if not n.is_finite() or n<0 or n>Decimal('9999999999.99'):raise ValueError
        return str(n.quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
    except (InvalidOperation,ValueError,TypeError):raise ValidationError(f'{label}: enter a non-negative amount with at most 10 integer digits')

def clean_rows(raw, company_id):
    if not isinstance(raw,list) or not raw or len(raw)>5000:raise ValidationError('Add between 1 and 5,000 employee rows')
    employees={e.pk:e for e in Employee.objects.filter(company_id=company_id)}
    seen=set();result=[]
    for row in raw:
        if not isinstance(row,dict):raise ValidationError('Invalid employee row')
        pk=row.get('employee_id')
        if type(pk)!=int or pk not in employees or pk in seen:raise ValidationError('Each employee must belong to this company and appear only once')
        seen.add(pk);emp=employees[pk]
        r={'employee_id':pk,'name':emp.full_name,'emp_no':emp.emp_no}
        for k in MONEY:r[k]=amount(row.get(k,'0'),f'{emp.emp_no}: {k}')
        basis=row.get('epf_basis','custom')
        if basis not in ('basic','custom'):raise ValidationError('Select a valid EPF calculation method')
        r['epf_basis']=basis
        if basis=='basic':r['epf_salary']=r['basic_salary']
        for k in OPTIONAL:r[k]=amount(row.get(k),f'{emp.emp_no}: {k}',True)
        r['epf_rate']=amount(row.get('epf_rate','8'),'EPF rate')
        if Decimal(r['epf_rate'])>100 or (r['format_percent'] is not None and Decimal(r['format_percent'])>100):raise ValidationError('Percentages must be between 0 and 100')
        r['notes']=str(row.get('notes',''))
        if len(r['notes'])>1000:raise ValidationError('Notes must be at most 1,000 characters')
        epf=(Decimal(r['epf_salary'])*Decimal(r['epf_rate'])/100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        r['epf']=str(epf)
        r['salary_amount']=str(Decimal(r['basic_salary'])-sum(Decimal(r[k]) for k in ('leave_deductions','welfare','salary_advance'))-epf)
        result.append(r)
    return result

def payload(s):
    return {'id':s.pk,'company':s.company_id,'company_name':s.company_name,'year':s.year,'month':s.month,'status':s.status,'revision':s.revision,'rows':s.rows,'template_config':s.template_config,'total':str(sum((Decimal(r['salary_amount']) for r in s.rows),Decimal(0)))}

@endpoint()
@require_http_methods(['GET','POST','DELETE'])
def sheets(request,pk=None):
    require_hr(request.user)
    qs=SalarySheet.objects.filter(company__in=companies(request.user))
    if request.method=='GET':
        if pk:return JsonResponse({'item':payload(get_object_or_404(qs,pk=pk))})
        page=max(1,int(request.GET.get('page',1)))
        return JsonResponse({'items':[{k:v for k,v in payload(s).items() if k!='rows'} for s in qs[(page-1)*50:page*50]],'count':qs.count()})
    data=body(request)
    if not isinstance(data,dict):raise ValidationError('Expected a form object')
    with transaction.atomic():
        if request.method=='DELETE' and not pk:raise ValidationError('Select a salary sheet')
        if pk:
            s=get_object_or_404(qs.select_for_update(),pk=pk)
            if s.status!='draft':raise ValidationError('Finalized sheets cannot be edited. Keep the original and record adjustments separately.')
            if data.get('revision')!=s.revision:return JsonResponse({'error':'This sheet changed in another window. Reload before saving.'},status=409)
            before=plain(s)
            if request.method=='DELETE':
                audit(request.user,'salary_sheet.delete',s,before);s.delete()
                return JsonResponse({'ok':True})
        else:
            form=SheetHeader(data)
            if not form.is_valid():raise ValidationError('Select a company, year (2000–2100) and month (1–12)')
            c=get_object_or_404(companies(request.user),pk=form.cleaned_data['company'])
            s=SalarySheet(company=c,company_name=c.name,year=form.cleaned_data['year'],month=form.cleaned_data['month']);before={}
        from . import template_payroll as template
        config=data.get('template_config',s.template_config)
        if s.pk and bool(config)!=bool(s.template_config):raise ValidationError('Existing sheet format cannot be changed')
        if config:
            s.template_config=template.clean_config(config,s.company)
            s.rows=template.clean_rows(data.get('rows'),s.company,s.template_config)
        else:s.rows=clean_rows(data.get('rows'),s.company_id)
        finalize=data.get('finalize',False)
        if type(finalize)!=bool:raise ValidationError('Invalid finalize option')
        if finalize:
            if any(Decimal(r['salary_amount'])<0 for r in s.rows):raise ValidationError('Resolve negative salary amounts before finalizing')
            s.status='finalized'
        if pk:s.revision+=1
        s.full_clean();s.save();audit(request.user,'salary_sheet.finalize' if finalize else 'salary_sheet.save',s,before)
        return JsonResponse({'item':payload(s)},status=200 if pk else 201)

@endpoint()
@require_http_methods(['GET'])
def options(request):
    require_hr(request.user)
    c=get_object_or_404(companies(request.user),pk=request.GET.get('company'))
    emps=Employee.objects.filter(company=c,status=Employee.ACTIVE).order_by('emp_no')
    # Suggestions only; the user reviews monthly amounts, including dated changes.
    from .template_payroll import default_config
    return JsonResponse({'template_config':default_config(c),'employees':[{'id':e.pk,'name':e.full_name,'emp_no':e.emp_no,'status':e.status,'basic_salary':str(e.basic_salary),'epf_eligible':e.epf_eligible,'epf_allowance':str(e.epf_allowance),'fixed_allowance':str(e.fixed_allowance)} for e in emps]})

def workbook_bytes(s):
    # Runtime export uses the project's existing openpyxl dependency. No employee
    # data from the reference workbook is bundled with this application.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.workbook.properties import CalcProperties
    w=Workbook();ws=w.active;ws.title=f'{calendar.month_abbr[s.month]}{str(s.year)[-2:]}'
    w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
    ws['B1']=safe(s.company_name);ws['B2']=f'{calendar.month_name[s.month]} {s.year} — {s.status.upper()}'
    ws['B1'].font=Font(size=14,bold=True)
    headers=['Name','Basic Salary','EPF Salary','Leave Deductions','Welfare','Salary Advance','EPF','Salary Amount','Notes','Appraisal','New Salary','Format %']
    for col,h in enumerate(headers,2):
        c=ws.cell(4,col,h);c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='D9EAD3');c.alignment=Alignment(wrap_text=True)
    for i,r in enumerate(s.rows,5):
        ws.cell(i,2,safe(r['name']))
        for col,k in enumerate(MONEY,3):ws.cell(i,col,Decimal(r[k]))
        ws.cell(i,8,f'=ROUND(D{i}*{r["epf_rate"]}/100,2)')
        ws.cell(i,9,f'=ROUND(C{i}-E{i}-F{i}-G{i}-H{i},2)')
        ws.cell(i,10,safe(r.get('notes','')))
        for col,k in enumerate(OPTIONAL,11):
            if r.get(k) is not None:ws.cell(i,col,Decimal(r[k])/(100 if k=='format_percent' else 1))
        for col in range(3,14):ws.cell(i,col).number_format='0.00%' if col==13 else '#,##0.00;[Red](#,##0.00)'
    end=4+len(s.rows);total=end+2;ws.cell(total,2,'Total')
    for col in ('C','D','E','F','G','H','I','K','L'):ws[f'{col}{total}']=f'=SUM({col}5:{col}{end})';ws[f'{col}{total}'].number_format='#,##0.00'
    for cell in ws[total]:cell.font=Font(bold=True)
    for row in ws.iter_rows(min_row=4,max_row=end,min_col=2,max_col=13):
        for c in row:c.border=Border(bottom=Side(style='hair',color='D9E2DC'))
    for col in 'CDEFGHIKL':ws.column_dimensions[col].width=19
    ws.column_dimensions['A'].width=3;ws.column_dimensions['B'].width=38;ws.column_dimensions['J'].width=35;ws.column_dimensions['M'].width=14
    ws.row_dimensions[4].height=32;ws.freeze_panes='C5';ws.auto_filter.ref=f'B4:M{end}'
    ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:4';ws.print_area=f'B1:M{total+2}'
    ws.cell(total+2,2,'Manual salary worksheet. Appraisal columns are reference only; no APIT/ETF/loan posting or payslips are generated.')
    b=io.BytesIO();w.save(b);return b.getvalue()

def safe(v):return "'"+v if v and v[0] in '=+-@' else v

@endpoint()
@require_http_methods(['GET'])
def export(request,pk):
    require_hr(request.user)
    s=get_object_or_404(SalarySheet.objects.filter(company__in=companies(request.user)),pk=pk)
    if s.template_config:
        from .template_payroll import workbook_bytes as template_bytes
        data=template_bytes(s)
    else:data=workbook_bytes(s)
    response=HttpResponse(data,content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="salary-sheet-{s.year}-{s.month:02}-{s.status}.xlsx"'
    response['Cache-Control']='no-store'
    return response
