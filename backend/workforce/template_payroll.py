"""Monthly component sheets using the uploaded salary template, with explicit payroll handoff."""
import io
import json
from pathlib import Path
from decimal import Decimal as D, ROUND_HALF_UP
from copy import deepcopy
from zipfile import ZipFile
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import JsonResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_http_methods
from core.models import Employee
from .api import endpoint, body, companies
from .models import SalarySheet
from .services import require_hr, audit, plain

LABELS=['Cost of living allowance','Food allowance','Other EPF-able allowance','Travelling allowance','Performance / incentive allowance','Other allowance']
FIELDS=['basic_salary']+[f'allowance_{i}' for i in range(6)]+['no_pay_days','ot_hours','bonus','salary_advance','welfare','loan_recovery','other_deductions']
HEADERS=['Emp No *','Employee name','Basic salary *']+LABELS+['No-pay days','OT hours','Bonus / incentive (not EPF)','Salary advance','Welfare','Loan recovery','Other deductions','Remarks','Gross pay','EPF/ETF earnings','EPF 8% (employee)','EPF 12% (employer)','ETF 3%','Net pay (before APIT)']
RESULTS=['gross_salary','epf_base','epf_employee','epf_employer','etf_employer','salary_amount']

def default_config(c):
    return {'version':2,'allowances':[{'name':n,'epf':i<3} for i,n in enumerate(LABELS)],'no_pay_divisor':str(c.no_pay_divisor),'ot_hourly_divisor':str(c.ot_hourly_divisor),'ot_multiplier':str(c.ot_multiplier)}

def clean_config(raw,c):
    out=default_config(c)
    items=raw.get('allowances') if isinstance(raw,dict) else None
    if not isinstance(items,list) or len(items)!=6:raise ValidationError('Six allowance columns are required')
    if D(out['no_pay_divisor'])<=0 or D(out['ot_hourly_divisor'])<=0:raise ValidationError('Company pay divisors must be positive')
    out['allowances']=[]
    for item in items:
        if not isinstance(item,dict) or type(item.get('epf'))!=bool:raise ValidationError('Choose Yes or No for each allowance')
        name=str(item.get('name','')).strip()
        if not name or len(name)>100:raise ValidationError('Allowance names must contain 1–100 characters')
        if type(item.get('enabled',True))!=bool:raise ValidationError('Invalid allowance enabled flag')
        out['allowances'].append({'name':name,'epf':item['epf'],'enabled':item.get('enabled',True)})
    custom=raw.get('custom_columns',[])
    if not isinstance(custom,list) or len(custom)>30:raise ValidationError('Maximum 30 custom columns')
    out['custom_columns']=[]
    seen=set()
    import re
    for col in custom:
        if not isinstance(col,dict):raise ValidationError('Invalid custom column')
        key=col.get('key','');name=str(col.get('name','')).strip()
        if not isinstance(key,str) or not re.fullmatch(r'custom_[a-zA-Z0-9_]{1,60}',key) or key in seen:raise ValidationError('Invalid or duplicate custom column key')
        if not name or len(name)>100:raise ValidationError('Column names must contain 1–100 characters')
        seen.add(key);out['custom_columns'].append({'key':key,'name':name})
    visible=raw.get('column_visibility',{})
    allowed={'emp_no','name','remarks',*FIELDS,*RESULTS,*seen}
    if not isinstance(visible,dict) or any(k not in allowed or type(v)!=bool for k,v in visible.items()):raise ValidationError('Column visibility must be Yes or No')
    removed=raw.get('removed_columns',[])
    if not isinstance(removed,list) or any(not isinstance(k,str) or k not in allowed for k in removed):raise ValidationError('Invalid removed columns')
    out['removed_columns']=list(dict.fromkeys(removed))
    names=raw.get('header_names',{})
    if not isinstance(names,dict) or any(k not in allowed or not isinstance(v,str) or not v.strip() or len(v)>100 for k,v in names.items()):raise ValidationError('Invalid column headers')
    out['header_names']={k:v.strip() for k,v in names.items()}
    out['column_visibility']=visible
    return out

def money(n):return D(n).quantize(D('.01'),rounding=ROUND_HALF_UP)

def totals(r,cfg):
    v={k:D(r[k]) for k in FIELDS}
    for i,a in enumerate(cfg['allowances']):
        if not a.get('enabled',True):v[f'allowance_{i}']=D('0')
    for k in cfg.get('removed_columns',[]):
        if k in v:v[k]=D('0')
    base=v['basic_salary'];fixed=base+sum(v[f'allowance_{i}'] for i in range(6))
    fraction=v['no_pay_days']/D(cfg['no_pay_divisor'])
    ot=base/D(cfg['ot_hourly_divisor'])*D(cfg['ot_multiplier'])*v['ot_hours']
    gross=money(fixed*(1-fraction)+ot+v['bonus'])
    epfbase=money((base+sum(v[f'allowance_{i}'] for i,a in enumerate(cfg['allowances']) if a['epf']))*(1-fraction))
    ee=money(epfbase*D('.08'));er=money(epfbase*D('.12'));etf=money(epfbase*D('.03'))
    net=gross-ee-sum(v[k] for k in ('salary_advance','welfare','loan_recovery','other_deductions'))
    return dict(zip(RESULTS,map(str,[gross,epfbase,ee,er,etf,money(net)])))

def clean_rows(rows,c,cfg):
    from .salary_sheets import amount
    if not isinstance(rows,list) or not 1<=len(rows)<=5000:raise ValidationError('Enter 1–5,000 employee rows')
    emps={e.pk:e for e in Employee.objects.filter(company=c)};seen=set();out=[]
    for row in rows:
        if not isinstance(row,dict):raise ValidationError('Invalid salary row')
        pk=row.get('employee_id')
        if type(pk)!=int or pk not in emps or pk in seen:raise ValidationError('Employee must belong to this company and appear once')
        seen.add(pk);e=emps[pk]
        r={'employee_id':pk,'name':e.full_name,'emp_no':e.emp_no}
        for k in FIELDS:r[k]=amount(row.get(k,'0'),f'{e.emp_no}: {k}')
        for i,a in enumerate(cfg['allowances']):
            if not a.get('enabled',True):r[f'allowance_{i}']='0.00'
        for k in cfg.get('removed_columns',[]):
            if k in FIELDS:r[k]='0.00'
        if D(r['no_pay_days'])>D(cfg['no_pay_divisor']):raise ValidationError('No-pay days exceed the company divisor')
        if D(r['ot_hours'])>D('9999.99'):raise ValidationError('OT hours exceed 9999.99')
        r['remarks']=str(row.get('remarks',''))
        if len(r['remarks'])>255:raise ValidationError('Remarks must be 255 characters or fewer')
        custom=row.get('custom_values',{})
        if not isinstance(custom,dict):raise ValidationError('Invalid custom column values')
        r['custom_values']={}
        for col in cfg.get('custom_columns',[]):
            value=str(custom.get(col['key'],''))
            if len(value)>255:raise ValidationError('Custom column values must be 255 characters or fewer')
            r['custom_values'][col['key']]=value
        if 'remarks' in cfg.get('removed_columns',[]):r['remarks']=''
        r.update(totals(r,cfg));out.append(r)
    return out

@endpoint()
@require_http_methods(['POST'])
def upload(request):
    """Validate and preview only. The browser merges these rows before explicit save."""
    require_hr(request.user)
    c=get_object_or_404(companies(request.user),pk=request.POST.get('company'))
    year=int(request.POST.get('year'));month=int(request.POST.get('month'))
    if not 2000<=year<=2100 or not 1<=month<=12:raise ValidationError('Invalid month')
    file=request.FILES.get('file')
    if not file or not file.name.lower().endswith('.xlsx') or file.size>5*1024*1024:raise ValidationError('Select an .xlsx file up to 5 MB')
    try:
        data=file.read()
        with ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist())>30*1024*1024:raise ValidationError('Workbook expands beyond 30 MB')
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        if 'Payroll' not in w.sheetnames:raise ValidationError('Missing Payroll worksheet')
        s=w['Payroll']
        if s.max_row>5008:raise ValidationError('Maximum 5,000 employee rows')
        if s['B2'].value not in (None,'',c.code):raise ValidationError('Company code does not match')
        if s['B3'].value not in (None,'',f'{year}-{month:02}'):raise ValidationError('Payroll month does not match')
        metadata={}
        if '_Column settings' in w.sheetnames:
            metadata=json.loads(w['_Column settings']['A1'].value)
            metadata=clean_config(metadata,c)
        headers=[s.cell(8,i).value for i in range(1,24)]
        keys=['emp_no','name',*FIELDS,'remarks',*RESULTS]
        if any(headers[i]!=metadata.get('header_names',{}).get(keys[i],HEADERS[i]) for i in list(range(3))+list(range(9,23))):raise ValidationError('Header row 8 does not match the salary template')
        cfg=metadata or default_config(c)
        for i in range(6):
            flag=s.cell(7,i+4).value
            if flag not in ('Yes','No'):raise ValidationError('Each allowance flag in row 7 must be Yes or No')
            cfg['allowances'][i]={**cfg['allowances'][i],'name':headers[i+3],'epf':flag=='Yes'}
        cfg=clean_config(cfg,c)
        emps={e.emp_no:e for e in Employee.objects.filter(company=c)};rows=[]
        for cells in s.iter_rows(min_row=9,max_col=17):
            values=[cell.value for cell in cells]
            if all(v in (None,'') for v in values):continue
            if any(cell.data_type=='f' for cell in cells):raise ValidationError('Input cells A:Q must contain values, not formulas')
            code=str(values[0]).strip() if values[0] is not None else ''
            if code not in emps:raise ValidationError(f'Unknown employee number: {code}')
            if values[2] in (None,''):raise ValidationError(f'{code}: Basic salary is required')
            r={'employee_id':emps[code].pk,'remarks':values[16] or ''}
            r.update({k:values[i+2] if values[i+2] is not None else '0' for i,k in enumerate(FIELDS)})
            rows.append(r)
        cleaned=clean_rows(rows,c,cfg);w.close()
    except ValidationError:raise
    except Exception:raise ValidationError('Cannot read this workbook. Use the salary upload template.')
    return JsonResponse({'rows':cleaned,'template_config':cfg})

def workbook_bytes(s):
    # Application runtime export, retaining the supplied workbook layout and formulas.
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    from .salary_sheets import safe
    w=load_workbook(Path(__file__).parent/'templates_data'/'salary-upload.xlsx');ws=w['Payroll']
    w['How to fill']['A13']='7. Save as .xlsx. In Payroll > Salary sheets select company/month, Upload salary Excel, review and save. Finalize and apply to open payroll when ready.'
    ws['B2']=s.company.code;ws['B3']=f'{s.year}-{s.month:02}'
    cfg=s.template_config
    ws['F2']=D(cfg['no_pay_divisor']);ws['F3']=D(cfg['ot_hourly_divisor']);ws['H3']=D(cfg['ot_multiplier'])
    ws['A4']=f'{s.company_name} — {s.status.upper()} — template checks before APIT; review statutory payroll separately.'
    for i,a in enumerate(cfg['allowances']):ws.cell(7,i+4,'Yes' if a['epf'] else 'No');ws.cell(8,i+4,safe(a['name']))
    end=max(108,len(s.rows)+8)
    for i,r in enumerate(s.rows,9):
        ws.cell(i,1,safe(r['emp_no']));ws.cell(i,2,safe(r['name']))
        for j,k in enumerate(FIELDS,3):ws.cell(i,j,D(r[k]))
        ws.cell(i,17,safe(r.get('remarks','')))
        if i>108:
            from copy import copy
            for col in range(1,24):ws.cell(i,col)._style=copy(ws.cell(9,col)._style)
        for col in range(18,24):
            ws.cell(i,col,Translator(ws.cell(9,col).value,origin=ws.cell(9,col).coordinate).translate_formula(ws.cell(i,col).coordinate))
    for col in range(18,24):
        from openpyxl.utils import get_column_letter
        letter=get_column_letter(col);ws.cell(4,col,f'=SUM({letter}9:{letter}{end})')
    from openpyxl.utils import get_column_letter
    keys=['emp_no','name',*FIELDS,'remarks',*RESULTS]
    for j,key in enumerate(keys,1):
        if key in cfg.get('header_names',{}):ws.cell(8,j,safe(cfg['header_names'][key]))
        ws.column_dimensions[get_column_letter(j)].hidden=key in cfg.get('removed_columns',[]) or not cfg.get('column_visibility',{}).get(key,True)
    for i,a in enumerate(cfg['allowances']):
        if not a.get('enabled',True):ws.column_dimensions[get_column_letter(i+4)].hidden=True
    for j,col in enumerate(cfg.get('custom_columns',[]),24):
        ws.cell(8,j,safe(col['name']))
        ws.column_dimensions[get_column_letter(j)].width=24
        ws.column_dimensions[get_column_letter(j)].hidden=not cfg.get('column_visibility',{}).get(col['key'],True)
        for i,r in enumerate(s.rows,9):ws.cell(i,j,safe(r.get('custom_values',{}).get(col['key'],'')))
    last=get_column_letter(23+len(cfg.get('custom_columns',[])))
    ws.print_area=f'A1:{last}{end}';ws.auto_filter.ref=f'A8:{last}{end}'
    settings=w.create_sheet('_Column settings');settings['A1']=json.dumps(cfg);settings.sheet_state='hidden'
    b=io.BytesIO();w.save(b);return b.getvalue()

@endpoint()
@require_http_methods(['POST'])
def apply_to_payroll(request,pk):
    require_hr(request.user)
    from payroll.models import start_period, PayrollPeriod, PayrollInput, eligible_employees
    from .models import Loan
    with transaction.atomic():
        s=get_object_or_404(SalarySheet.objects.select_for_update().filter(company__in=companies(request.user)),pk=pk)
        if s.status!='finalized' or not s.template_config:raise ValidationError('Finalize a new template sheet first')
        period,_=start_period(s.company,s.year,s.month,carry_forward=False)
        period=PayrollPeriod.objects.select_for_update().get(pk=period.pk)
        if period.status!='open':raise ValidationError('Payroll must be open. Approved or locked payroll cannot be overwritten.')
        allowed=set(eligible_employees(s.company,s.year,s.month).values_list('pk',flat=True))
        if any(r['employee_id'] not in allowed for r in s.rows):raise ValidationError('Sheet includes employees outside this payroll month. Review employment dates.')
        if any(str(getattr(s.company,k))!=str(s.template_config[k]) for k in ('no_pay_divisor','ot_hourly_divisor','ot_multiplier')):raise ValidationError('Company calculation settings changed. Review the sheet before transfer.')
        for r in s.rows:
            if D(r['loan_recovery']) and Loan.objects.filter(employee_id=r['employee_id'],status='approved').exists():raise ValidationError('Manual loan recovery overlaps an approved loan. Resolve before transfer to avoid double deductions.')
            inp,_=PayrollInput.objects.get_or_create(period=period,employee_id=r['employee_id']);before=plain(inp)
            inp.salary_template={'sheet_id':s.pk,'config':deepcopy(s.template_config),'basic_salary':r['basic_salary'],'allowances':[{**a,'amount':r[f'allowance_{i}']} for i,a in enumerate(s.template_config['allowances'])]}
            for k in ('no_pay_days','ot_hours','bonus','salary_advance','welfare','loan_recovery','other_deductions','remarks'):setattr(inp,k,r[k])
            inp.epf_earnings=0;inp.other_earnings=0
            inp.apit_override=None;inp.tax_note=''
            inp.save();audit(request.user,'payroll.template_import',inp,before)
        audit(request.user,'salary_sheet.apply',s)
    return JsonResponse({'period_id':period.pk,'message':'Inputs transferred. Review APIT and automatic loan deductions before processing payroll.'})
